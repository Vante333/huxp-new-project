"""导出 PA10 已完成需求评估表 Excel（附件3格式）。"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "PA10-requirements-assessment-202510-202602.xlsx"
JIRA_BROWSE_BASE = "https://dragonpass.atlassian.net/browse/"

HEADERS = ["需求说明", "ticket", "风险等级", "流量等级", "需求等级", "备注", "上线日期"]

EXCLUDE_KEYS = {"PA10-112", "PA10-125", "PA10-135"}

# 需求说明, ticket, 风险, 流量, 等级
ROWS: list[tuple[str, str, str, str, str]] = [
    (
        "【1.0】礼宾车创建订单/订单详情接口新增境外联系方式字段 【全部API版本】",
        "PA10-73",
        "M",
        "M",
        "B",
    ),
    (
        "【1.0API】休息室资源详情展示setmeal套餐列表 【全部API版本】",
        "PA10-77",
        "L",
        "M",
        "C",
    ),
    (
        "飞常准出行项目，CIP资源详情接口",
        "PA10-82",
        "M",
        "M",
        "B",
    ),
    (
        "【SaaS 1.0】新增回调重试机制 【全部API版本】",
        "PA10-84",
        "M",
        "M",
        "B",
    ),
    (
        "针对礼宾车品类，优化权益服务订单同步逻辑（equity_use_record）",
        "PA10-92",
        "M",
        "M",
        "B",
    ),
    (
        "兴业白金卡境外接送机项目 -下单接口支持境外礼宾车订单，修改时间检验规则",
        "PA10-94",
        "M",
        "H",
        "B",
    ),
    (
        "飞猪辅营 - 新增饭票码查询接口",
        "PA10-95",
        "M",
        "M",
        "B",
    ),
    (
        "站点列表接口去重",
        "PA10-97",
        "L",
        "M",
        "C",
    ),
    (
        "礼宾车权益MQ，再增加一个推送场景，订单完成后再触发权益使用MQ",
        "PA10-115",
        "M",
        "L",
        "B",
    ),
    (
        "切换Google API key",
        "PA10-117",
        "L",
        "L",
        "C",
    ),
    (
        "品类推送查漏补缺",
        "PA10-119",
        "L",
        "L",
        "C",
    ),
    (
        "【文档优化】鉴权文档",
        "PA10-127",
        "L",
        "L",
        "C",
    ),
    (
        "回调支持贵宾MG品类",
        "PA10-130",
        "M",
        "L",
        "B",
    ),
    (
        "新增嘉宾资源的详情查询接口",
        "PA10-132",
        "M",
        "M",
        "B",
    ),
    (
        "回调新增品类类型字段",
        "PA10-138",
        "L",
        "L",
        "C",
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")
LEGEND_FILL = PatternFill("solid", fgColor="F2F2F2")
REMARK = "无需测试"


def _LoadDotenv() -> None:
    env_path = Path(__file__).resolve().parent.parent / ".env"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                k, v = k.strip(), v.strip()
                if k and not os.environ.get(k):
                    os.environ[k] = v


def TicketUrl(ticket: str) -> str:
    key = ticket.strip()
    if key.startswith("http"):
        return key
    return f"{JIRA_BROWSE_BASE}{key}"


def _FetchGoLiveDates(keys: list[str]) -> dict[str, str]:
    """优先 Jira Start date（customfield_10015），否则用进入 Done 类目日期。"""
    _LoadDotenv()
    base = (os.environ.get("JIRA_BASE_URL") or "https://dragonpass.atlassian.net").rstrip("/")
    auth = (os.environ["JIRA_EMAIL"], os.environ["JIRA_API_TOKEN"])
    headers = {"Accept": "application/json"}
    out: dict[str, str] = {}
    chunk_size = 50
    for i in range(0, len(keys), chunk_size):
        chunk = keys[i : i + chunk_size]
        jql = "key in (" + ",".join(chunk) + ")"
        token: str | None = None
        while True:
            body: dict = {
                "jql": jql,
                "maxResults": 100,
                "fields": ["customfield_10015", "statuscategorychangedate", "resolutiondate"],
            }
            if token:
                body["nextPageToken"] = token
            r = requests.post(
                f"{base}/rest/api/3/search/jql",
                auth=auth,
                headers={**headers, "Content-Type": "application/json"},
                json=body,
                timeout=60,
            )
            r.raise_for_status()
            data = r.json()
            for it in data.get("issues", []):
                f = it.get("fields") or {}
                key = it.get("key", "")
                start_date = f.get("customfield_10015") or ""
                if start_date:
                    out[key] = str(start_date)[:10]
                    continue
                rd = f.get("resolutiondate") or ""
                if rd:
                    out[key] = rd[:10]
                    continue
                changed = f.get("statuscategorychangedate") or ""
                out[key] = changed[:10] if changed else ""
            token = data.get("nextPageToken")
            if data.get("isLast") or not token:
                break
    return out


def main() -> None:
    rows = [r for r in ROWS if r[1] not in EXCLUDE_KEYS]
    keys = [r[1] for r in rows]
    try:
        go_live_dates = _FetchGoLiveDates(keys)
    except Exception as exc:
        print(f"WARN: Jira go-live date fetch failed: {exc}", file=sys.stderr)
        go_live_dates = {}

    wb = Workbook()
    ws = wb.active
    ws.title = "PA10需求评估"

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row = 2
    for desc, ticket, risk, traffic, level in rows:
        ws.cell(row=row, column=1, value=desc)
        url = TicketUrl(ticket)
        ticket_cell = ws.cell(row=row, column=2, value=url)
        ticket_cell.hyperlink = url
        ticket_cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=3, value=risk)
        ws.cell(row=row, column=4, value=traffic)
        ws.cell(row=row, column=5, value=level)
        ws.cell(row=row, column=6, value=REMARK)
        ws.cell(row=row, column=7, value=go_live_dates.get(ticket, ""))
        row += 1

    legends = [
        ("需求等级说明", ""),
        ("A级别", "100%投入"),
        ("B级别", "关注修改点，回归核心功能"),
        ("C级别", "开发自测"),
        ("", ""),
        ("风险/流量 H", "高"),
        ("风险/流量 M", "中"),
        ("风险/流量 L", "低"),
    ]
    for i, (k, v) in enumerate(legends, 1):
        ws.cell(row=i, column=8, value=k).font = Font(bold=(i == 1))
        ws.cell(row=i, column=9, value=v)
        if i == 1:
            ws.cell(row=i, column=8).fill = LEGEND_FILL

    widths = [55, 52, 10, 10, 10, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 28

    ws.freeze_panes = "A2"
    wb.save(OUT)
    print(f"Wrote {OUT} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
