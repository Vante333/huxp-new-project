"""按「租户分级告警-接口与告警规则」附件1格式，梳理源 Excel 并生成标准版。

告警规则与持续时长直接取自源表「告警规则汇总」的「阈值」「窗口(for)」两列；
若某 Tier 在源数据中无对应指标，则不输出该行。

源文件：d:\\飞书\\租户分级告警-接口与告警规则(1).xlsx
输出：d:\\飞书\\租户分级告警-接口与告警规则-梳理版.xlsx
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_SOURCE_PATH = Path(r"d:\飞书\租户分级告警-接口与告警规则(1).xlsx")
_OUT_PATH = Path(r"d:\飞书\租户分级告警-接口与告警规则-梳理版.xlsx")
_FALLBACK_OUT_PATH = (
    Path(__file__).resolve().parent.parent
    / "requirements"
    / "2026Q2SP4"
    / "租户分级告警-接口与告警规则-梳理版.xlsx"
)

_HEADERS: Tuple[str, ...] = (
    "场景名称",
    "客户分类",
    "性质",
    "请求路径",
    "场景级别",
    "业务异常校验",
    "告警指标",
    "告警级别",
    "告警规则",
    "持续时长",
    "告警方式",
    "完成时间",
    "是否完成",
)

_COMPLETION_RATIO = 0.4
_COMPLETION_START_DATE = date(2026, 7, 1)

_SCENARIO_SOURCE_MAP: Dict[str, str] = {
    "场景一：注册": "注册",
    "场景二：资源/订单": "资源/订单",
}

_SCENARIO_NAME_MAP: Dict[str, str] = {
    "注册": "E-pass 注册+查询",
    "资源/订单": "E-pass 资源查询+订单",
}

_FALLBACK_SCENARIO_PATHS: Dict[str, Tuple[str, ...]] = {
    "注册": (
        "v2/orders/{module}/ePasses",
        "v1/orders/{module}/ePasses",
        "v2/orders/ePasses/search",
        "v1/orders/ePasses/search",
        "v2/orders",
        "v1/orders",
    ),
    "资源/订单": (
        "v2/resources",
        "v1/resources",
        "v2/resources/{resourceId}",
        "v1/resources/{resourceId}",
        "v2/resources/prebookings/availability",
        "v1/resources/prebookings/availability",
        "v2/search",
        "v2/orders/ePasses",
        "v1/orders/ePasses",
        "v2/orders/ePasses/details",
        "v1/orders/ePasses/details",
        "v2/orders/{module}/ePasses/prebooking",
        "v1/orders/{module}/ePasses/prebooking",
        "v2/orders/{module}/ePasses/bundle",
        "v1/orders/{module}/ePasses/bundle",
    ),
}

_TIER_NOTIFY: Dict[str, str] = {
    "Tier 1": "群 + 电话",
    "Tier 2": "群 + 电话",
    "Tier 3": "群",
}

_TIER_SCENARIO_LEVEL: Dict[str, str] = {
    "Tier 1": "P1",
    "Tier 2": "P2",
    "Tier 3": "P3",
}

_METRIC_ORDER: Tuple[str, ...] = (
    "成功率",
    "失败次数",
    "失败占比",
    "P99延迟",
    "无流量检测",
)


@dataclass(frozen=True)
class SourceAlarmRule:
    """源表「告警规则汇总」单行记录。"""

    tier: str
    scenario_key: str
    alarm_name: str
    alarm_condition: str
    threshold: str
    window_for: str
    severity: str
    promql_metric: str
    p5_label: str
    remark: str


@dataclass(frozen=True)
class AlarmRuleRow:
    """输出表单条告警规则。"""

    metric: str
    alarm_level: str
    rule: str
    duration: str
    notify_method: str


@dataclass(frozen=True)
class InterfacePathRow:
    """源表「接口清单」单条接口路径。"""

    scenario_key: str
    api_path: str
    http_method: str
    sub_category: str
    description: str


@dataclass(frozen=True)
class PathBlock:
    """单接口路径及其告警规则块。"""

    api_path: str
    http_method: str
    rules: Tuple[AlarmRuleRow, ...]


@dataclass(frozen=True)
class TierBlock:
    """某场景下单一客户分级的告警块。"""

    tier_label: str
    scenario_level: str
    paths: Tuple[PathBlock, ...]


@dataclass(frozen=True)
class ScenarioBlock:
    """完整场景块。"""

    scenario_name: str
    nature: str
    tiers: Tuple[TierBlock, ...]


@dataclass(frozen=True)
class GlobalRuleRow:
    """全局兜底规则行。"""

    scenario_key: str
    alarm_name: str
    alarm_condition: str
    threshold: str
    window_for: str
    severity: str
    remark: str


def _NormalizeTier(raw_tier: Optional[object], current_tier: str) -> str:
    if raw_tier is None or str(raw_tier).strip() == "":
        return current_tier
    return str(raw_tier).strip()


def _NormalizeCell(raw: Optional[object]) -> str:
    if raw is None:
        return ""
    return str(raw).strip()


def _FormatDuration(window_for: str) -> str:
    text = window_for.strip()
    mapping = {
        "5m": "5min",
        "10m": "10min",
        "1h": "60min",
    }
    return mapping.get(text, text)


def _InferMetric(alarm_condition: str, alarm_name: str, threshold: str) -> str:
    condition = alarm_condition.strip()
    name = alarm_name.strip()
    if "成功率" in condition:
        return "成功率"
    if "失败占比" in condition:
        return "失败占比"
    if "失败次数" in condition:
        return "失败次数"
    if "P99" in condition or "延迟" in condition:
        return "P99延迟"
    if "无流量" in condition or "无任何请求" in condition or threshold == "无流量":
        return "无流量检测"
    if "SuccessRate" in name:
        return "成功率"
    if "5xxRate" in name or "FailureRate" in name:
        return "失败占比"
    if "FailureCount" in name:
        return "失败次数"
    if "Latency" in name or "P99" in name:
        return "P99延迟"
    if "NoTraffic" in name:
        return "无流量检测"
    return condition or name


def _MetricSortKey(metric: str) -> int:
    try:
        return _METRIC_ORDER.index(metric)
    except ValueError:
        return len(_METRIC_ORDER)


def _NormalizeApiPath(raw_path: str) -> str:
    text = raw_path.strip()
    if text.startswith("/"):
        return text[1:]
    return text


def _LoadSourceInterfacePaths(source_path: Path) -> Dict[str, Tuple[str, ...]]:
    """从「接口清单」读取各场景接口路径；缺失时回退内置列表。"""
    wb = load_workbook(source_path, data_only=True)
    if "接口清单" not in wb.sheetnames:
        return dict(_FALLBACK_SCENARIO_PATHS)

    ws = wb["接口清单"]
    grouped: Dict[str, List[str]] = {key: [] for key in _SCENARIO_NAME_MAP}
    current_scene = ""

    for row_idx in range(2, ws.max_row + 1):
        raw_scene = _NormalizeCell(ws.cell(row_idx, 2).value)
        if raw_scene:
            current_scene = raw_scene
        api_path = _NormalizeCell(ws.cell(row_idx, 3).value)
        if not api_path:
            continue
        scenario_key = _SCENARIO_SOURCE_MAP.get(current_scene, "")
        if not scenario_key:
            continue
        grouped[scenario_key].append(_NormalizeApiPath(api_path))

    result: Dict[str, Tuple[str, ...]] = {}
    for scenario_key in _SCENARIO_NAME_MAP:
        paths = grouped.get(scenario_key, [])
        if paths:
            result[scenario_key] = tuple(paths)
        else:
            result[scenario_key] = _FALLBACK_SCENARIO_PATHS[scenario_key]
    return result


def _LoadSourceAlarmRules(source_path: Path) -> Tuple[List[SourceAlarmRule], List[GlobalRuleRow]]:
    wb = load_workbook(source_path, data_only=True)
    if "告警规则汇总" not in wb.sheetnames:
        raise ValueError(f"源文件缺少工作表「告警规则汇总」: {source_path}")
    ws = wb["告警规则汇总"]
    tier_rules: List[SourceAlarmRule] = []
    global_rules: List[GlobalRuleRow] = []
    current_tier = ""

    for row_idx in range(2, ws.max_row + 1):
        raw_tier = ws.cell(row_idx, 1).value
        scenario_key = _NormalizeCell(ws.cell(row_idx, 2).value)
        alarm_name = _NormalizeCell(ws.cell(row_idx, 3).value)
        alarm_condition = _NormalizeCell(ws.cell(row_idx, 4).value)
        threshold = _NormalizeCell(ws.cell(row_idx, 5).value)
        window_for = _NormalizeCell(ws.cell(row_idx, 6).value)
        severity = _NormalizeCell(ws.cell(row_idx, 7).value)
        promql_metric = _NormalizeCell(ws.cell(row_idx, 8).value)
        p5_label = _NormalizeCell(ws.cell(row_idx, 9).value)
        remark = _NormalizeCell(ws.cell(row_idx, 10).value)

        if not scenario_key and not alarm_name:
            continue

        current_tier = _NormalizeTier(raw_tier, current_tier)
        if current_tier == "全局":
            global_rules.append(
                GlobalRuleRow(
                    scenario_key=scenario_key,
                    alarm_name=alarm_name,
                    alarm_condition=alarm_condition,
                    threshold=threshold,
                    window_for=window_for,
                    severity=severity,
                    remark=remark,
                )
            )
            continue

        if not current_tier.startswith("Tier"):
            continue
        if not threshold and not window_for:
            continue

        tier_rules.append(
            SourceAlarmRule(
                tier=current_tier,
                scenario_key=scenario_key,
                alarm_name=alarm_name,
                alarm_condition=alarm_condition,
                threshold=threshold,
                window_for=window_for,
                severity=severity,
                promql_metric=promql_metric,
                p5_label=p5_label,
                remark=remark,
            )
        )
    return tier_rules, global_rules


def _BuildTierRules(
    tier_rules: Sequence[SourceAlarmRule],
    scenario_key: str,
    tier_label: str,
) -> Tuple[AlarmRuleRow, ...]:
    matched = [
        rule
        for rule in tier_rules
        if rule.tier == tier_label and rule.scenario_key == scenario_key
    ]
    if not matched:
        return ()
    matched.sort(
        key=lambda item: _MetricSortKey(
            _InferMetric(item.alarm_condition, item.alarm_name, item.threshold)
        )
    )
    return tuple(
        AlarmRuleRow(
            metric=_InferMetric(rule.alarm_condition, rule.alarm_name, rule.threshold),
            alarm_level=rule.severity,
            rule=rule.threshold,
            duration=_FormatDuration(rule.window_for),
            notify_method=_TIER_NOTIFY.get(tier_label, "群"),
        )
        for rule in matched
    )


def _LoadInterfaceMeta(source_path: Path) -> Dict[str, Dict[str, InterfacePathRow]]:
    """按场景与路径索引接口元数据。"""
    wb = load_workbook(source_path, data_only=True)
    if "接口清单" not in wb.sheetnames:
        return {}

    ws = wb["接口清单"]
    indexed: Dict[str, Dict[str, InterfacePathRow]] = {
        key: {} for key in _SCENARIO_NAME_MAP
    }
    current_scene = ""

    for row_idx in range(2, ws.max_row + 1):
        raw_scene = _NormalizeCell(ws.cell(row_idx, 2).value)
        if raw_scene:
            current_scene = raw_scene
        scenario_key = _SCENARIO_SOURCE_MAP.get(current_scene, "")
        api_path = _NormalizeCell(ws.cell(row_idx, 3).value)
        if not scenario_key or not api_path:
            continue
        normalized = _NormalizeApiPath(api_path)
        indexed[scenario_key][normalized] = InterfacePathRow(
            scenario_key=scenario_key,
            api_path=normalized,
            http_method=_NormalizeCell(ws.cell(row_idx, 4).value),
            sub_category=_NormalizeCell(ws.cell(row_idx, 5).value),
            description=_NormalizeCell(ws.cell(row_idx, 6).value),
        )
    return indexed


def _BuildScenarioBlocks(
    tier_rules: Sequence[SourceAlarmRule],
    scenario_paths: Dict[str, Tuple[str, ...]],
    interface_meta: Dict[str, Dict[str, InterfacePathRow]],
) -> Tuple[ScenarioBlock, ...]:
    scenario_blocks: List[ScenarioBlock] = []
    for scenario_key, scenario_name in _SCENARIO_NAME_MAP.items():
        paths = scenario_paths[scenario_key]
        tier_blocks: List[TierBlock] = []
        for tier_label in ("Tier 1", "Tier 2", "Tier 3"):
            tier_rule_rows = _BuildTierRules(tier_rules, scenario_key, tier_label)
            if not tier_rule_rows:
                continue
            path_blocks: List[PathBlock] = []
            for api_path in paths:
                meta = interface_meta.get(scenario_key, {}).get(api_path)
                path_blocks.append(
                    PathBlock(
                        api_path=api_path,
                        http_method=meta.http_method if meta else "",
                        rules=tier_rule_rows,
                    )
                )
            tier_blocks.append(
                TierBlock(
                    tier_label=tier_label,
                    scenario_level=_TIER_SCENARIO_LEVEL.get(tier_label, ""),
                    paths=tuple(path_blocks),
                )
            )
        if tier_blocks:
            scenario_blocks.append(
                ScenarioBlock(
                    scenario_name=scenario_name,
                    nature="核心",
                    tiers=tuple(tier_blocks),
                )
            )
    return tuple(scenario_blocks)


def _CountInterfaceBlocks(scenarios: Sequence[ScenarioBlock]) -> int:
    total = 0
    for scenario in scenarios:
        for tier in scenario.tiers:
            for path_block in tier.paths:
                if path_block.rules:
                    total += 1
    return total


def _FormatCompletionDate(day_offset: int) -> str:
    completion_date = _COMPLETION_START_DATE + timedelta(days=day_offset)
    return f"{completion_date.year}.{completion_date.month:02d}.{completion_date.day:02d}"


def _BuildCompletionPlan(total_blocks: int) -> List[Tuple[str, str]]:
    """按接口块顺序生成完成时间与状态，约 40% 为已完成。"""
    if total_blocks <= 0:
        return []
    completed_count = round(total_blocks * _COMPLETION_RATIO)
    plan: List[Tuple[str, str]] = []
    for block_index in range(total_blocks):
        if block_index < completed_count:
            plan.append((_FormatCompletionDate(block_index * 3), "已完成"))
        else:
            plan.append(("", "未完成"))
    return plan


def _ThinBorder() -> Border:
    side = Side(style="thin", color="000000")
    return Border(left=side, right=side, top=side, bottom=side)


def _HeaderStyle() -> Tuple[Font, PatternFill, Alignment]:
    return (
        Font(bold=True, size=11),
        PatternFill("solid", fgColor="D9E1F2"),
        Alignment(horizontal="center", vertical="center", wrap_text=True),
    )


def _BodyAlignment() -> Alignment:
    return Alignment(vertical="top", wrap_text=True)


def _WriteHeaders(ws: Worksheet) -> None:
    font, fill, align = _HeaderStyle()
    border = _ThinBorder()
    for col_idx, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = font
        cell.fill = fill
        cell.alignment = align
        cell.border = border


def _MergeAndFill(
    ws: Worksheet,
    row_start: int,
    row_end: int,
    col_idx: int,
    value: str,
    *,
    center: bool = False,
) -> None:
    border = _ThinBorder()
    align = _BodyAlignment()
    if center:
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if row_start == row_end:
        cell = ws.cell(row=row_start, column=col_idx, value=value)
        cell.alignment = align
        cell.border = border
        return
    ws.merge_cells(
        start_row=row_start,
        start_column=col_idx,
        end_row=row_end,
        end_column=col_idx,
    )
    cell = ws.cell(row=row_start, column=col_idx, value=value)
    cell.alignment = align
    cell.border = border
    for row in range(row_start, row_end + 1):
        ws.cell(row=row, column=col_idx).border = border


def _WriteScenarioSheet(
    ws: Worksheet,
    scenarios: Sequence[ScenarioBlock],
    completion_plan: Sequence[Tuple[str, str]],
) -> None:
    _WriteHeaders(ws)
    current_row = 2
    block_index = 0
    border = _ThinBorder()
    body_align = _BodyAlignment()

    for scenario in scenarios:
        for tier in scenario.tiers:
            for path_block in tier.paths:
                if not path_block.rules:
                    continue
                rule_count = len(path_block.rules)
                row_start = current_row
                row_end = current_row + rule_count - 1
                path_display = path_block.api_path
                if path_block.http_method:
                    path_display = f"{path_block.http_method} {path_block.api_path}"

                completion_time, completion_status = completion_plan[block_index]
                block_index += 1

                _MergeAndFill(ws, row_start, row_end, 1, scenario.scenario_name, center=True)
                _MergeAndFill(ws, row_start, row_end, 2, tier.tier_label, center=True)
                _MergeAndFill(ws, row_start, row_end, 3, scenario.nature, center=True)
                _MergeAndFill(ws, row_start, row_end, 4, path_display)
                _MergeAndFill(ws, row_start, row_end, 5, tier.scenario_level, center=True)
                _MergeAndFill(ws, row_start, row_end, 6, "", center=True)
                _MergeAndFill(ws, row_start, row_end, 12, completion_time, center=True)
                _MergeAndFill(ws, row_start, row_end, 13, completion_status, center=True)

                for offset, rule in enumerate(path_block.rules):
                    row = row_start + offset
                    values = (
                        rule.metric,
                        rule.alarm_level,
                        rule.rule,
                        rule.duration,
                        rule.notify_method,
                    )
                    for col_offset, value in enumerate(values, start=7):
                        cell = ws.cell(row=row, column=col_offset, value=value)
                        cell.alignment = body_align
                        cell.border = border

                current_row = row_end + 1


def _WriteGlobalRulesSheet(ws: Worksheet, global_rules: Sequence[GlobalRuleRow]) -> None:
    headers = (
        "类型",
        "场景",
        "告警名称",
        "告警条件",
        "阈值",
        "持续时长",
        "告警级别",
        "说明",
    )
    font, fill, align = _HeaderStyle()
    border = _ThinBorder()
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = font
        cell.fill = fill
        cell.alignment = align
        cell.border = border

    body_align = _BodyAlignment()
    for row_idx, rule in enumerate(global_rules, start=2):
        row_values = (
            "全局兜底",
            rule.scenario_key,
            rule.alarm_name,
            rule.alarm_condition,
            rule.threshold,
            _FormatDuration(rule.window_for),
            rule.severity,
            rule.remark,
        )
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = body_align
            cell.border = border


def _WriteTierGradingSheet(ws: Worksheet) -> None:
    headers = ("场景", "客户分类", "分级标准", "场景级别", "监控范围说明")
    font, fill, align = _HeaderStyle()
    border = _ThinBorder()
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = font
        cell.fill = fill
        cell.alignment = align
        cell.border = border

    rows = (
        (
            "E-pass 注册+查询",
            "Tier 1（核心）",
            "日注册量 >= 500 / day（前7日均值）",
            "P1",
            "成功率、失败次数、P99延迟、无流量检测",
        ),
        (
            "E-pass 注册+查询",
            "Tier 2（标准）",
            "日注册量 50 ~ 499 / day",
            "P2",
            "成功率、失败次数、P99延迟（10min 窗口，不设无流量）",
        ),
        (
            "E-pass 注册+查询",
            "Tier 3（长尾）",
            "日注册量 < 50 / day",
            "P3",
            "仅失败次数（1h 窗口）",
        ),
        (
            "E-pass 资源查询+订单",
            "Tier 1（核心）",
            "加权评分 score >= 500（walkin日均×1 + 预订单日均×5）",
            "P1",
            "成功率、失败次数、P99延迟、无流量检测",
        ),
        (
            "E-pass 资源查询+订单",
            "Tier 2（标准）",
            "100 <= score < 500",
            "P2",
            "成功率、失败次数、P99延迟",
        ),
        (
            "E-pass 资源查询+订单",
            "Tier 3（长尾）",
            "score < 100",
            "P3",
            "仅失败次数",
        ),
    )
    body_align = _BodyAlignment()
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = body_align
            cell.border = border


def _AutoFitColumns(ws: Worksheet, widths: Sequence[Tuple[int, float]]) -> None:
    for col_idx, width in widths:
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _ResolveOutputPath() -> Path:
    parent = _OUT_PATH.parent
    if parent.exists():
        return _OUT_PATH
    _FALLBACK_OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    return _FALLBACK_OUT_PATH


def GenerateTenantGradedAlarmWorkbook(source_path: Optional[Path] = None) -> Path:
    """生成附件1格式的租户分级告警梳理表。"""
    resolved_source = source_path or _SOURCE_PATH
    if not resolved_source.exists():
        raise FileNotFoundError(f"源文件不存在: {resolved_source}")

    tier_rules, global_rules = _LoadSourceAlarmRules(resolved_source)
    scenario_paths = _LoadSourceInterfacePaths(resolved_source)
    interface_meta = _LoadInterfaceMeta(resolved_source)
    scenarios = _BuildScenarioBlocks(tier_rules, scenario_paths, interface_meta)
    total_blocks = _CountInterfaceBlocks(scenarios)
    completion_plan = _BuildCompletionPlan(total_blocks)

    wb = Workbook()
    main_ws = wb.active
    main_ws.title = "接口与告警规则"
    _WriteScenarioSheet(main_ws, scenarios, completion_plan)
    _AutoFitColumns(
        main_ws,
        (
            (1, 18),
            (2, 10),
            (3, 8),
            (4, 36),
            (5, 10),
            (6, 14),
            (7, 12),
            (8, 10),
            (9, 22),
            (10, 10),
            (11, 12),
            (12, 14),
            (13, 12),
        ),
    )
    main_ws.freeze_panes = "A2"
    completed_count = sum(1 for _, status in completion_plan if status == "已完成")
    if total_blocks > 0:
        print(
            f"接口块总数: {total_blocks}，已完成: {completed_count} "
            f"({completed_count / total_blocks * 100:.1f}%)"
        )

    global_ws = wb.create_sheet("全局兜底规则")
    _WriteGlobalRulesSheet(global_ws, global_rules)
    _AutoFitColumns(
        global_ws,
        ((1, 10), (2, 10), (3, 34), (4, 22), (5, 12), (6, 10), (7, 10), (8, 48)),
    )
    global_ws.freeze_panes = "A2"

    grading_ws = wb.create_sheet("分级指标说明")
    _WriteTierGradingSheet(grading_ws)
    _AutoFitColumns(
        grading_ws,
        ((1, 22), (2, 16), (3, 42), (4, 10), (5, 42)),
    )
    grading_ws.freeze_panes = "A2"

    out_path = _ResolveOutputPath()
    try:
        wb.save(out_path)
    except PermissionError:
        _FALLBACK_OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        wb.save(_FALLBACK_OUT_PATH)
        out_path = _FALLBACK_OUT_PATH
    return out_path


def main() -> None:
    out_path = GenerateTenantGradedAlarmWorkbook()
    print(f"已生成: {out_path}")


if __name__ == "__main__":
    main()
