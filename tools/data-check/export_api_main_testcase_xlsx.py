"""从 Apifox MCP 多分支 OpenAPI 合并导出接口测试用例 Excel 模板。

默认以 **API main**（BranchId 309401）为主集，并补充 **API Q1SP4+5**（309343）、
**Q2 local offer**（309401，与 main 同分支）中 main 未覆盖的接口。

用法（PowerShell，仓库根目录）::

    python tools/data-check/export_api_main_testcase_xlsx.py

快照目录：``tech-docs/2026Q2_SP3/apifox-oas-snapshots/``（可用 MCP refresh 后手动更新）。

模块分类与 ``{module}`` 赋值对齐 ``api-testcases/docs/api-main-call-chains.md``；
含 ``{module}`` 的 path 按品类展开；全部用例合并为单一工作表输出。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_TOOLS_DIR = Path(__file__).resolve().parent
if str(_TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(_TOOLS_DIR))
from api_main_endpoint_testcase_hints import GetCheckpointAndExpected
from api_main_call_chain_catalog import BuildCatalogEntries, EndpointCatalogEntry

JsonDict = Dict[str, Any]
ExpandedKey = Tuple[str, str]  # concrete_path, method
TestcaseRow = Tuple[str, str, str, str, str, str, str, str, str, str, str, str]

DEFAULT_AUTOMATION_STATUS = "已补充"
PRIORITY_COL_LETTER = "F"
AUTOMATION_COL_LETTER = "H"
SUMMARY_SHEET_NAME = "统计"
DATA_SHEET_NAME = "接口用例"

HEADERS: Sequence[str] = (
    "模块*",
    "功能点*",
    "检查点*",
    "详细说明*",
    "预期结果*",
    "优先级*",
    "平台*",
    "接口自动化",
    "UI自动化",
    "接口自动化脚本",
    "提交人*",
    "备注",
)

HEADER_FILL = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DEFAULT_OUT = Path("testcases/2026Q2_SP3/API-main-接口测试用例模板.xlsx")
SNAPSHOT_DIR = Path("tech-docs/2026Q2_SP3/apifox-oas-snapshots")

DEFAULT_SOURCES: Sequence[Tuple[str, str, str]] = (
    ("API main", "api-main.json", "309401"),
    ("Q2 local offer", "q2-local-offer.json", "309401"),
    ("API Q1SP4+5", "q1sp45.json", "309343"),
)

_HTTP_METHODS = frozenset({"get", "post", "put", "patch", "delete", "head", "options"})


def _RepoRoot() -> Path:
    return Path(__file__).resolve().parents[2]


def _LoadOas(path: Path) -> JsonDict:
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as exc:
        raise RuntimeError(f"无法读取 OAS 文件: {path}") from exc
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"OAS JSON 解析失败: {path}") from exc
    if not isinstance(data, dict):
        raise RuntimeError(f"OAS 根节点应为 dict: {path}")
    return data


def _SecondSegmentIsCore(path: str, scenario: str) -> bool:
    if scenario.startswith(("Membership", "EPass", "ESIM")):
        return True
    segments = [s.lower() for s in path.strip("/").split("/") if s]
    if len(segments) < 2:
        return False
    core = {"auth", "memberships", "entitlements", "orders", "users", "internal"}
    return segments[1] in core


def _EnhanceCheckpointForModule(checkpoint: str, entry: EndpointCatalogEntry) -> str:
    if entry.module_slug and entry.module_int is not None:
        return (
            f"{checkpoint}（pathModule={entry.module_slug}，body.module={entry.module_int}）"
        )
    return checkpoint


def _BuildRow(
    entry: EndpointCatalogEntry,
    source_label: str,
    branch_id: str,
    *,
    supplement: bool,
) -> TestcaseRow:
    path = entry.concrete_path
    method = entry.method
    summary = entry.summary
    module = entry.BuildModuleColumn()
    function_point = summary or f"{method} {path}"
    checkpoint, expected = GetCheckpointAndExpected(path, method, summary)
    checkpoint = _EnhanceCheckpointForModule(checkpoint, entry)
    detail = f"{method} {path}{entry.BuildDetailSuffix()}"
    if summary:
        detail = f"{detail} — {summary}"
    priority = "P0" if _SecondSegmentIsCore(path, entry.scenario) else "P1"
    platform = "接口"
    automation = DEFAULT_AUTOMATION_STATUS
    remark = f"主链路={entry.scenario}；Apifox {source_label} BranchId={branch_id}"
    if entry.template_path != entry.concrete_path:
        remark = f"{remark}；模板={entry.template_path}"
    if supplement:
        remark = f"{remark}；补充自其他分支（main 未包含）"
    return (
        module,
        function_point,
        checkpoint,
        detail,
        expected,
        priority,
        platform,
        automation,
        "",
        "",
        "",
        remark,
    )


def _ExtractEndpointRows(
    oas: JsonDict,
    source_label: str,
    branch_id: str,
    *,
    supplement: bool,
) -> List[Tuple[ExpandedKey, TestcaseRow, str, int]]:
    paths = oas.get("paths")
    if not isinstance(paths, dict):
        raise RuntimeError(f"{source_label} OAS 缺少 paths 节点")
    items: List[Tuple[ExpandedKey, TestcaseRow, str, int]] = []
    for path in sorted(paths.keys(), key=_SortPathKey):
        path_item = paths[path]
        if not isinstance(path_item, dict):
            continue
        for key, value in path_item.items():
            if key.lower() not in _HTTP_METHODS:
                continue
            if not isinstance(value, dict):
                continue
            method = key.upper()
            summary = str(value.get("summary") or "").strip()
            for entry in BuildCatalogEntries(path, method, summary):
                row = _BuildRow(
                    entry,
                    source_label,
                    branch_id,
                    supplement=supplement,
                )
                expanded_key: ExpandedKey = (entry.concrete_path, entry.method)
                items.append((expanded_key, row, entry.top_group, entry.sort_key[0]))
    return items


def _SortPathKey(path: str) -> Tuple[int, str]:
    return (len(path.strip("/").split("/")), path)


def _MergeSources(
    sources: Sequence[Tuple[str, Path, str]],
) -> Tuple[List[TestcaseRow], Dict[str, int]]:
    merged: Dict[ExpandedKey, Tuple[TestcaseRow, int]] = {}
    stats = {"main": 0, "supplement_q1": 0, "supplement_q2": 0, "skipped_dup": 0, "expanded": 0}
    primary_label = sources[0][0] if sources else "API main"

    for idx, (label, oas_path, branch_id) in enumerate(sources):
        if not oas_path.is_file():
            raise RuntimeError(f"缺少 OAS 快照: {oas_path}（请先用 MCP refresh 落盘）")
        oas = _LoadOas(oas_path)
        is_primary = idx == 0
        for key, row, _top_group, scenario_order in _ExtractEndpointRows(
            oas,
            label,
            branch_id,
            supplement=not is_primary,
        ):
            if key in merged:
                stats["skipped_dup"] += 1
                continue
            merged[key] = (row, scenario_order)
            stats["expanded"] += 1
            if is_primary:
                stats["main"] += 1
            elif label == "API Q1SP4+5":
                stats["supplement_q1"] += 1
            elif label == "Q2 local offer":
                stats["supplement_q2"] += 1

    if not merged:
        raise RuntimeError(f"合并后无接口；请检查 {primary_label} 快照")

    ordered_rows = [
        row for _key, (row, _order) in sorted(
            merged.items(),
            key=lambda item: (item[1][1], item[0][0], item[0][1]),
        )
    ]
    return ordered_rows, stats


def _ApplyHeaderStyle(ws: Worksheet) -> None:
    for col_idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28


def _AutoColumnWidth(ws: Worksheet, max_width: int = 48) -> None:
    for col_idx in range(1, len(HEADERS) + 1):
        letter = get_column_letter(col_idx)
        max_len = len(str(HEADERS[col_idx - 1]))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is None:
                    continue
                for line in str(cell.value).splitlines():
                    max_len = max(max_len, len(line))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), max_width)


def _WriteSheet(ws: Worksheet, rows: Sequence[TestcaseRow]) -> None:
    _ApplyHeaderStyle(ws)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    _AutoColumnWidth(ws)


def _QuoteSheetName(name: str) -> str:
    escaped = name.replace("'", "''")
    return f"'{escaped}'"


def _WriteSummarySheet(wb: Workbook) -> None:
    ws = wb.create_sheet(title=SUMMARY_SHEET_NAME, index=0)
    title_font = Font(bold=True, size=14)
    header_fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    data_sheet = _QuoteSheetName(DATA_SHEET_NAME)

    ws["A1"] = "接口自动化完成率统计"
    ws["A1"].font = title_font
    ws.merge_cells("A1:D1")

    headers = ("维度", "用例总数", "已补充数", "完成率")
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows_spec = (
        ("P0", "P0"),
        ("P1", "P1"),
        ("合计", None),
    )
    for row_idx, (label, priority) in enumerate(rows_spec, start=4):
        ws.cell(row=row_idx, column=1, value=label)
        total_cell = ws.cell(row=row_idx, column=2)
        done_cell = ws.cell(row=row_idx, column=3)
        rate_cell = ws.cell(row=row_idx, column=4)
        b_ref = f"B{row_idx}"
        c_ref = f"C{row_idx}"
        if priority is None:
            total_cell.value = "=B4+B5"
            done_cell.value = "=C4+C5"
        else:
            total_cell.value = (
                f"=COUNTIF({data_sheet}!${PRIORITY_COL_LETTER}:${PRIORITY_COL_LETTER},\"{priority}\")"
            )
            done_cell.value = (
                f"=COUNTIFS({data_sheet}!${PRIORITY_COL_LETTER}:${PRIORITY_COL_LETTER},\"{priority}\","
                f"{data_sheet}!${AUTOMATION_COL_LETTER}:${AUTOMATION_COL_LETTER},"
                f"\"{DEFAULT_AUTOMATION_STATUS}\")"
            )
        rate_cell.value = f'=IF({b_ref}=0,"N/A",{c_ref}/{b_ref})'
        rate_cell.number_format = "0.00%"

    ws.cell(row=8, column=1, value="说明").font = Font(bold=True)
    ws.cell(
        row=9,
        column=1,
        value=(
            f"「接口自动化」列默认「{DEFAULT_AUTOMATION_STATUS}」；"
            f"修改「{DATA_SHEET_NAME}」sheet 中该列后，本表完成率自动重算。"
        ),
    )
    ws.merge_cells("A9:D9")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12


def _WriteWorkbook(rows: Sequence[TestcaseRow], out_path: Path) -> int:
    wb = Workbook()
    default_ws = wb.active
    default_ws.title = DATA_SHEET_NAME
    _WriteSheet(default_ws, rows)
    if not rows:
        raise RuntimeError("无可用工作表数据")
    _WriteSummarySheet(wb)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(rows)


def Main() -> int:
    root = _RepoRoot()
    parser = argparse.ArgumentParser(description="从多分支 Apifox OAS 合并导出接口测试用例 Excel")
    parser.add_argument(
        "--snapshot-dir",
        type=Path,
        default=root / SNAPSHOT_DIR,
        help="OAS 快照目录",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=root / DEFAULT_OUT,
        help="输出 xlsx 路径",
    )
    args = parser.parse_args()
    snapshot_dir = args.snapshot_dir.resolve()
    sources: List[Tuple[str, Path, str]] = [
        (label, snapshot_dir / filename, branch_id)
        for label, filename, branch_id in DEFAULT_SOURCES
    ]
    rows, stats = _MergeSources(sources)
    total = _WriteWorkbook(rows, args.out.resolve())

    p0_total = sum(1 for row in rows if row[5] == "P0")
    p1_total = sum(1 for row in rows if row[5] == "P1")
    p0_done = sum(1 for row in rows if row[5] == "P0" and row[7] == DEFAULT_AUTOMATION_STATUS)
    p1_done = sum(1 for row in rows if row[5] == "P1" and row[7] == DEFAULT_AUTOMATION_STATUS)
    all_done = p0_done + p1_done
    pct = lambda done, cnt: (f"{done / cnt * 100:.2f}%" if cnt else "N/A")

    print(
        f"已导出 {total} 条接口用例 -> {args.out.resolve()}\n"
        f"  工作表: {SUMMARY_SHEET_NAME} + {DATA_SHEET_NAME}（合并单 sheet）\n"
        f"  完成率（接口自动化默认「{DEFAULT_AUTOMATION_STATUS}」）:\n"
        f"    P0: {p0_done}/{p0_total} = {pct(p0_done, p0_total)}\n"
        f"    P1: {p1_done}/{p1_total} = {pct(p1_done, p1_total)}\n"
        f"    合计: {all_done}/{total} = {pct(all_done, total)}\n"
        f"  API main 原始: {stats['main']} 条\n"
        f"  补充 Q1SP4+5: {stats['supplement_q1']} 条\n"
        f"  补充 Q2 local offer: {stats['supplement_q2']} 条\n"
        f"  {{module}} 展开后合计: {stats['expanded']} 条\n"
        f"  跳过重複: {stats['skipped_dup']} 条"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(Main())
