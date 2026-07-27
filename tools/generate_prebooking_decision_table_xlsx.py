"""生成订单流预收费/权益支付/现金抵扣判定组合表 Excel。"""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_REPO_ROOT = Path(__file__).resolve().parent.parent
_OUT_PATH = _REPO_ROOT / "requirements" / "2026Q2SP4" / "订单流-判定组合表-预收费权益现金.xlsx"
_FALLBACK_OUT_PATH = _REPO_ROOT / "requirements" / "2026Q2SP4" / "订单流-判定组合表-预收费权益现金-新版.xlsx"

_HEADERS = [
    "规则",
    "prebooking fee 付费方式",
    "prebooking fee 价格",
    "超人次费用",
    "超人次费用价格",
    "客户预收费预期",
    "权益支付预期",
    "现金（权益抵扣）预期",
]


def _JoinLines(lines: Sequence[str]) -> str:
    return "\n".join(lines)


def _CustomerOrCashExpected(prebooking_lines: Sequence[str], over_fee_lines: Sequence[str]) -> str:
    return _JoinLines(
        [
            "（1）预约费用：",
            *prebooking_lines,
            "（2）超人次费用：",
            *over_fee_lines,
            "（3）收银台：",
            "应付价格非0时拉取收银台；应付价格=0时不拉取收银台",
        ]
    )


def _EntitlementExpected(lines: Sequence[str]) -> str:
    return _JoinLines(
        [
            "走权益支付商品：",
            *lines,
        ]
    )


_CUSTOMER_PREBOOKING_LINES = [
    "prebooking fee 由客户承担，预约费用=0",
]

_USER_PREBOOKING_NON_ZERO_LINES = [
    "prebooking fee 由用户承担",
    "预约费用=预约单价*3",
]

_USER_PREBOOKING_ZERO_LINES = [
    "prebooking fee 由用户承担，但预约单价=0",
    "预约费用=0",
]

_OVER_FEE_NON_ZERO_LINES = [
    "a. 使用权益点数0时，超人次费用=3*超人次价格",
    "b. 使用权益点数1时，超人次费用=2*超人次价格",
    "应付价格=预约费用+超人次费用",
]

_OVER_FEE_NOT_SET_LINES = [
    "未设置超人次费用，超人次费用=0",
    "应付价格=预约费用",
]

_OVER_FEE_ZERO_LINES = [
    "超人次价格=0，超人次费用=0",
    "不产生权益抵扣金额",
    "应付价格=预约费用",
]

_ENTITLEMENT_COMMON_LINES = [
    "按权益点数扣减",
    "权益足够则下单成功",
    "权益不足则返回权益不足",
    "不展示现金金额信息，payment={}",
]

_ROWS: list[list[str]] = [
    [
        "R1",
        "客户",
        "非0",
        "设置",
        "非0",
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NON_ZERO_LINES),
        _EntitlementExpected(_ENTITLEMENT_COMMON_LINES),
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NON_ZERO_LINES),
    ],
    [
        "R2",
        "客户",
        "0",
        "设置",
        "非0",
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NON_ZERO_LINES),
        _EntitlementExpected(_ENTITLEMENT_COMMON_LINES),
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NON_ZERO_LINES),
    ],
    [
        "R3",
        "客户",
        "非0",
        "不设置",
        "-",
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NOT_SET_LINES),
        _EntitlementExpected(["无现金金额；权益足够则下单成功", "不展示现金金额信息，payment={}"]),
        _CustomerOrCashExpected(_CUSTOMER_PREBOOKING_LINES, _OVER_FEE_NOT_SET_LINES),
    ],
    [
        "R4",
        "用户",
        "非0",
        "设置",
        "非0",
        _CustomerOrCashExpected(_USER_PREBOOKING_NON_ZERO_LINES, _OVER_FEE_NON_ZERO_LINES),
        _EntitlementExpected(_ENTITLEMENT_COMMON_LINES),
        _CustomerOrCashExpected(_USER_PREBOOKING_NON_ZERO_LINES, _OVER_FEE_NON_ZERO_LINES),
    ],
    [
        "R5",
        "用户",
        "0",
        "设置",
        "非0",
        _CustomerOrCashExpected(_USER_PREBOOKING_ZERO_LINES, _OVER_FEE_NON_ZERO_LINES),
        _EntitlementExpected(_ENTITLEMENT_COMMON_LINES),
        _CustomerOrCashExpected(_USER_PREBOOKING_ZERO_LINES, _OVER_FEE_NON_ZERO_LINES),
    ],
    [
        "R6",
        "用户",
        "非0",
        "设置",
        "0",
        _CustomerOrCashExpected(_USER_PREBOOKING_NON_ZERO_LINES, _OVER_FEE_ZERO_LINES),
        _EntitlementExpected(["无超人次金额", "权益足够则下单成功", "不展示现金金额信息，payment={}"]),
        _CustomerOrCashExpected(_USER_PREBOOKING_NON_ZERO_LINES, _OVER_FEE_ZERO_LINES),
    ],
    [
        "R7",
        "未配置",
        "-",
        "未配置",
        "-",
        _JoinLines(["未配置客户预收费商品时：", "预期返回「不支持预收款」或商品不可用类错误"]),
        _JoinLines(["未配置权益支付商品时：", "预期返回「不支持权益支付」或权益商品不可用类错误"]),
        _JoinLines(["未配置现金支付商品时：", "预期返回「商品未配置现金支付，不支持拉起收银台」"]),
    ],
]

_RULES: list[str] = [
    "计算规则",
    "应付价格 = 超人次费用 + 预约费用",
    "",
    "预约费用",
    "· prebooking fee 付费方式=客户：预约费用=0",
    "· prebooking fee 付费方式=用户：预约费用=预约单价*3",
    "",
    "超人次费用（预约3人）",
    "· 使用权益点数0：超人次费用=3*超人次价格",
    "· 使用权益点数1：超人次费用=2*超人次价格",
    "",
    "收银台规则",
    "· 只要应付价格不为0，都要拉取收银台",
    "· 应付价格=0，不拉取收银台",
    "",
    "补充",
    "· 客户预收费和现金（权益抵扣）的应付价格预期类似",
    "· 现金（权益抵扣）商品如果不支持权益抵扣，则不按使用权益点数抵扣超人次费用，可按使用权益点数0的结果校验",
]


def _ApplyHeaderStyle(ws: Worksheet, row: int, col_count: int) -> None:
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="B4B4B4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def _ApplyBodyStyle(ws: Worksheet, start_row: int, end_row: int, col_count: int) -> None:
    thin = Side(style="thin", color="B4B4B4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(start_row, end_row + 1):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        ws.row_dimensions[row].height = 140


def _AutoColumnWidth(ws: Worksheet, col_count: int, max_width: int = 48) -> None:
    for col in range(1, col_count + 1):
        letter = get_column_letter(col)
        if col <= 5:
            ws.column_dimensions[letter].width = 13
        else:
            ws.column_dimensions[letter].width = max_width


def GenerateExcel() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "判定组合表"

    ws.append(_HEADERS)
    for row in _ROWS:
        ws.append(row)

    _ApplyHeaderStyle(ws, 1, len(_HEADERS))
    _ApplyBodyStyle(ws, 2, len(_ROWS) + 1, len(_HEADERS))
    _AutoColumnWidth(ws, len(_HEADERS), max_width=54)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"

    ws_rules = wb.create_sheet("计算规则说明")
    for line in _RULES:
        ws_rules.append([line])
    ws_rules.column_dimensions["A"].width = 80
    for row in ws_rules.iter_rows(min_row=1, max_row=ws_rules.max_row, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    _OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(_OUT_PATH)
        return _OUT_PATH
    except PermissionError:
        wb.save(_FALLBACK_OUT_PATH)
        return _FALLBACK_OUT_PATH


if __name__ == "__main__":
    path = GenerateExcel()
    print(f"已生成: {path}")
