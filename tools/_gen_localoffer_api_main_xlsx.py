# -*- coding: utf-8 -*-
"""按 API-main-接口测试用例模板.xlsx 格式生成 Local Offer 接口清单用例表。"""
from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

template = Path(r"d:\interface\interface_cursor\testcases\2026Q2_SP3\API-main-接口测试用例模板.xlsx")
out = Path(r"d:\interface\interface_cursor\testcases\2026Q2_SP3\LocalOffer-接口清单-测试用例表.xlsx")


def _RowValues(ws, row_idx: int) -> list:
    return [ws.cell(row_idx, c).value for c in range(1, 13)]



def Main() -> None:
    shutil.copy2(template, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["接口用例"]

    source_wb = openpyxl.load_workbook(template, data_only=True)
    source_ws = source_wb["接口用例"]
    epass_create = _RowValues(source_ws, 50)
    epass_bundle = _RowValues(source_ws, 63)
    epass_prebooking = _RowValues(source_ws, 58)

    style_src = 2
    saved_styles = []
    for col in range(1, 13):
        src = ws.cell(style_src, col)
        saved_styles.append((copy(src._style), src.number_format, src.alignment.horizontal))

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    new_rows = [
        [
            "Membership-Extensions扩展-",
            "Update Membership Extensions",
            "携带有效 Token 与 X-Request-ID，对已注册 membershipId 更新 extensions 非标字段（UPSERT/DELETE）",
            "PATCH /v2/memberships/extensions — Update Membership Extensions",
            "返回 HTTP 200，code=0；extensions 字段更新成功",
            "P1",
            "接口",
            "已覆盖",
            None,
            None,
            None,
            "脑图路径=Membership-Extensions扩展-；Apifox API main BranchId=309401；PAPI-2548",
        ],
        [
            "公共能力-ProjectResources项目资源-",
            "RetrieveProjectResourcesList",
            "携带有效 Token，传入合法 projectId 与单个 module 查询项目可用资源列表",
            "GET /api/v2/projects/{projectId}/resources — RetrieveProjectResourcesList",
            "返回 HTTP 200，code=0；返回指定 module 下项目可用资源列表（分页）",
            "P1",
            "接口",
            "未覆盖",
            None,
            None,
            None,
            "脑图路径=公共能力-ProjectResources项目资源-；Apifox API main BranchId=309401；PAPI-2553",
        ],
        [
            "公共能力-ProjectResources项目资源-",
            "RetrieveProjectResourcesList",
            "携带有效 Token，传入合法 projectId 与单个 module 查询项目可用资源列表",
            "GET /v2/projects/{projectId}/resources — RetrieveProjectResourcesList",
            "返回 HTTP 200，code=0；返回指定 module 下项目可用资源列表（分页）",
            "P1",
            "接口",
            "未覆盖",
            None,
            None,
            None,
            "脑图路径=公共能力-ProjectResources项目资源-；Apifox API main BranchId=309401；PAPI-2553；与 /api/v2 路径等价口径待联调确认",
        ],
        [
            "公共能力-ProgramResources批量查询-",
            "BatchRetrieveProgramResources",
            "携带有效 Token，传入合法 programIds 与单个 module 批量查询各 program 可用资源",
            "POST /v2/programs/resources/batch — BatchRetrieveProgramResources",
            "返回 HTTP 200，code=0；批量返回各 program 可用资源列表",
            "P1",
            "接口",
            "未覆盖",
            None,
            None,
            None,
            "脑图路径=公共能力-ProgramResources批量查询-；Apifox API main BranchId=309401；PAPI-2553 / PAPI-2620",
        ],
        [
            "公共能力-ProjectTransportHub项目交通枢纽-",
            "RetrieveProjectTransportHubList",
            "携带有效 Token，传入合法 projectId 与单个 module 查询项目可用交通枢纽列表",
            "GET /v2/projects/{projectId}/transportHubs — RetrieveProjectTransportHubList",
            "返回 HTTP 200，code=0；返回指定 module 下项目可用交通枢纽列表",
            "P1",
            "接口",
            "未覆盖",
            None,
            None,
            None,
            "脑图路径=公共能力-ProjectTransportHub项目交通枢纽-；Apifox API main BranchId=309401；PAPI-2553",
        ],
        [
            "公共能力-ProgramTransportHub批量查询-",
            "BatchRetrieveProgramTransportHubs",
            "携带有效 Token，传入合法 programIds 与单个 module 批量查询各 program 可用交通枢纽",
            "POST /v2/programs/transportHubs/batch — BatchRetrieveProgramTransportHubs",
            "返回 HTTP 200，code=0；批量返回各 program 可用交通枢纽列表",
            "P1",
            "接口",
            "未覆盖",
            None,
            None,
            None,
            "脑图路径=公共能力-ProgramTransportHub批量查询-；Apifox API main BranchId=309401；PAPI-2553",
        ],
        epass_create,
        epass_bundle,
        epass_prebooking,
    ]

    for idx, row in enumerate(new_rows, start=2):
        for col, val in enumerate(row, start=1):
            cell = ws.cell(idx, col, val)
            style, number_format, horizontal = saved_styles[col - 1]
            cell._style = copy(style)
            cell.number_format = number_format
            cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Rows: {len(new_rows)}")


if __name__ == "__main__":
    Main()
